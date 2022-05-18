# -*- coding: utf-8 -*-

__author__ = "joe-tester"


import os

from openpyxl.reader.excel import load_workbook
import xlsxwriter
import numpy as np
from matplotlib import pyplot as plt
import time

from common.configpath import TIME_PATH, CPU_MEM_PATH, ScreenShot_DIR
from utils.logger import logger, LOG


# The two-dimensional diagram supports displaying Chinese
plt.rcParams['font.sans-serif'] = ['Microsoft YaHei']
# time format
time_now = time.strftime("%Y%m%d%H%M%S", time.localtime())


@logger('save time of start app')
def start_app(times, start):
    try:
        workbook = xlsxwriter.Workbook(TIME_PATH)
        worksheet = workbook.add_worksheet('time')
        bold = workbook.add_format({'bold': 1})
        headings = ['启动次数', '启动时间']
        data = [times, start]
        worksheet.write_row('A1', headings, bold)
        worksheet.write_column('A2', data[0])
        worksheet.write_column('B2', data[1])
        chart1 = workbook.add_chart({'type': 'scatter',
                                     'subtype': 'straight_with_markers'})
        chart1.add_series({
            'name': '=time!$B$1',
            'categories': '=time!$A$2:$A$%s' % (len(start) + 1),
            'values': '=time!$B$2:$B$%s' % (len(start) + 1),
        })
        chart1.set_title({'name': '启动监测'})
        chart1.set_x_axis({'name': "启动次数"})
        chart1.set_y_axis({'name': '花费时间:ms'})
        chart1.set_style(11)
        worksheet.insert_chart('D2', chart1, {'x_offset': 25, 'y_offset': 10})
        workbook.close()
        LOG.info('save success!!!')
    except:
        LOG.info('save failed，because is: %s' % Exception)


@logger('save cpuinfo meminfo netflow_info batt_info')
def get_cpu(times, start_cpu, recv_list, send_list, total_list, mem_list, batt_list):
    try:
        # create workbook
        workbook = xlsxwriter.Workbook(CPU_MEM_PATH)
        # create sheet
        worksheet = workbook.add_worksheet('cpu')
        worksheet_netflow = workbook.add_worksheet('netflow')
        worksheet_mem = workbook.add_worksheet('mem')
        worksheet_batt = workbook.add_worksheet('batt')
        # set style
        bold = workbook.add_format({'bold': 1})
        # header
        headings = ['监控次数', 'cpu占用率%']
        headings_netflow = ['监控次数', '上行流量', '下行流量', '流量总计']
        headings_mem = ['监控次数', 'Pass占百分比']
        headings_batt = ["监控次数", '耗电百分比']
        # datas
        data_cpu = [times, start_cpu]
        data_netflow = [times, recv_list, send_list, total_list]
        data_mem = [times, mem_list]
        data_batt = [times, batt_list]
        # write cpuinfo to excel
        worksheet.write_row('A1', headings, bold)
        worksheet.write_column('A2', data_cpu[0])
        worksheet.write_column('B2', data_cpu[1])
        # write batt to excel
        worksheet_batt.write_row('A1', headings_batt, bold)
        worksheet_batt.write_column('A2', data_batt[0])
        worksheet_batt.write_column('B2', data_batt[1])
        # write netflow to excel
        worksheet_netflow.write_row('A1', headings_netflow, bold)
        worksheet_netflow.write_column('A2', data_netflow[0])
        worksheet_netflow.write_column('B2', data_netflow[2])
        worksheet_netflow.write_column('C2', data_netflow[1])
        worksheet_netflow.write_column('D2', data_netflow[3])
        # write meminfo to excel
        worksheet_mem.write_row('A1', headings_mem, bold)
        worksheet_mem.write_column('A2', data_mem[0])
        worksheet_mem.write_column('B2', data_mem[1])

        # Generate 2D map
        chart1 = workbook.add_chart({'type': 'scatter',
                                     'subtype': 'straight_with_markers'})
        chart2 = workbook.add_chart({'type': 'scatter',
                                     'subtype': 'straight_with_markers'})
        chart3 = workbook.add_chart({'type': 'scatter',
                                     'subtype': 'straight_with_markers'})
        chart4 = workbook.add_chart({'type': 'scatter',
                                     'subtype': 'straight_with_markers'})

        chart1.add_series({
            'name': '=cpu!$B$1',
            'categories': '=cpu!$A$2:$A$%s' % (len(times) + 1),
            'values': '=cpu!$B$2:$B$%s' % (len(times) + 1),
        })
        chart2.add_series({
            'name': '=netflow!$B$1',
            'categories': '=netflow!$A$2:$A$%s' % (len(times) + 1),
            'values': '=netflow!$B$2:$B$%s' % (len(times) + 1),

        })
        chart2.add_series({
            'name': '=netflow!$C$1',    # netflow
            'categories': '=netflow!$A$2:$A$%s' % (len(times) + 1),
            'values': '=netflow!$C$2:$C$%s' % (len(times) + 1),
        })
        chart2.add_series({
            'name': '=netflow!$D$1',
            'categories': '=netflow!$A$2:$A$%s' % (len(times) + 1),
            'values': '=netflow!$D$2:$D$%s' % (len(times) + 1),
        })
        chart3.add_series({
            'name': '=mem!$B$1',
            'categories': '=mem!$A$2:$A$%s' % (len(times) + 1),
            'values': '=mem!$B$2:$B$%s' % (len(times) + 1),
        })
        chart4.add_series({
            'name': '=batt!$B$1',
            'categories': '=batt!$A$2:$A$%s' % (len(times) + 1),
            'values': '=batt!$B$2:$B$%s' % (len(times) + 1),
        })

        chart1.set_title({'name': '进程cpu占用率'})
        chart1.set_x_axis({'name': "次数"})
        chart1.set_y_axis({'name': '占用:%'})
        chart1.set_style(11)

        chart2.set_title({'name': '流量统计曲线'})
        chart2.set_x_axis({'name': '次数'})
        chart2.set_y_axis({'name': '流量：k'})
        chart2.set_style(11)

        chart3.set_title({'name': '进程mem占有率'})
        chart3.set_x_axis({'name': '次数'})
        chart3.set_y_axis({'name': 'pass值：%'})
        chart3.set_style(11)

        chart4.set_title({'name': '手机剩余电量比'})
        chart4.set_x_axis({'name': "次数"})
        chart4.set_y_axis({'name': '电量:%'})
        chart4.set_style(11)

        worksheet_mem.insert_chart('F2', chart3, {'x_offset': 60, 'y_offset': 60})
        worksheet_netflow.insert_chart('F2', chart2, {'x_offset': 60, 'y_offset': 60})
        worksheet.insert_chart('D2', chart1, {'x_offset': 60, 'y_offset': 60})
        worksheet_batt.insert_chart('D2', chart4, {'x_offset': 60, 'y_offset': 60})

        workbook.close()
        LOG.info('Successfully saved collected data')
    except:
        LOG.info('Failed to save collected data: %s' % Exception)


class HandleExcel(object):
    def __init__(self, filename=CPU_MEM_PATH):
        self.filename = filename
        self.wb = load_workbook(self.filename)
        # self.ws = self.wb[self.sheetname] if self.sheetname is not None else self.wb.active
        # title
        # self.sheet_head_tuple = tuple(self.ws.iter_rows(max_row=self.ws.min_row, values_only=True))[0]
        self.times = 0
        self.cpus = []
        self.mems = []
        self.netflows = []
        self.uploads = []
        self.downloads = []
        self.batts = []
        # self.Cases = namedtuple("cases", self.sheet_head_tuple)

    def get_cpus(self):
        self.ws = self.wb["cpu"]
        self.times = self.ws.max_row - 1
        for tuple_data in self.ws.iter_rows(min_row=self.ws.min_row + 1, values_only=True):
            self.cpus.append(tuple_data[1])
            # self.cases_list.append(self.Cases(*tuple_data))
        return self.cpus

    def get_mems(self):
        self.ws = self.wb["mem"]
        self.times = self.ws.max_row - 1
        for tuple_data in self.ws.iter_rows(min_row=self.ws.min_row + 1, values_only=True):
            self.mems.append(tuple_data[1])
        return self.mems

    def get_netflows(self):
        self.ws = self.wb["netflow"]
        self.times = self.ws.max_row - 1
        for tuple_data in self.ws.iter_rows(min_row=self.ws.min_row + 1, values_only=True):
            self.netflows.append(tuple_data[3])
            self.uploads.append(tuple_data[1])
            self.downloads.append(tuple_data[2])
        return self.uploads, self.downloads, self.netflows

    def get_batts(self):
        self.ws = self.wb["batt"]
        self.times = self.ws.max_row - 1
        for tuple_data in self.ws.iter_rows(min_row=self.ws.min_row + 1, values_only=True):
            self.batts.append(tuple_data[1])
            # self.cases_list.append(self.Cases(*tuple_data))
        return self.batts


excel = HandleExcel()


def mapping_cpu(packagename):
    cpus = excel.get_cpus()
    cpus = list(map(float, cpus))
    times = [i for i in range(1, excel.times + 1)]
    xpoint = np.array(times)
    # Draw graphics based on data
    plt.figure(figsize=(11, 7), dpi=600)
    # Generate grid
    plt.grid(axis="y")
    
    # print("cpu列表：{}".format(cpus))

    # sum
    # total = 0
    # for value in cpus:
    #    total += value
    
    # average max min
    # average = round(total / len(cpus), 2)
    # cpu_h = sorted(cpus)
    # print("cpu平均值：{}".format(average))
    # print("cpu最低值：{}".format(cpu_h[0]))
    # print("cpu最高值：{}".format(cpu_h[len(cpu_h) - 1]))

    ypoint = np.array(cpus)

    plt.plot(xpoint, ypoint, color="green", linestyle="-", linewidth=1, label=packagename)

    # Axis range
    plt.ylim(min(ypoint) - 20, max(ypoint) + 20)
    plt.xlim(-1, len(xpoint) + 2)

    plt.xlabel('采集频率(次)', fontsize=16)
    plt.ylabel("cpu占用率%", fontsize=16)
    plt.title("APP进程CPU占用率%", fontsize=24)

    plt.legend()

    # Abscissa display interval
    # if len(times) <= 15:
    #     pass
    # else:
    #     t = int(len(times) / 15)
    #     plt.xticks(range(0, len(times), t))

    # plt.show()

    # If it is a date format, it needs to be rotated
    # plt.gcf().autofmt_xdate()

    path = os.path.join(ScreenShot_DIR, 'cpu_' + time_now)
    plt.savefig(path)
    
    
def mapping_mem(packagename):
    mems = excel.get_mems()
    mems = list(map(float, mems))
    times = [i for i in range(1, excel.times + 1)]
    # Draw graphics based on data
    plt.figure(figsize=(11, 7), dpi=600)
    # Generate grid
    plt.grid(axis="y")
    
    xpoint = np.array(times)
    # print("内存值：{}".format(mems))
    ypoint = np.array(mems)
    # start draw
    plt.plot(xpoint, ypoint, color="blue", linestyle="-", linewidth=1, label=packagename)
    # Axis range
    plt.ylim(min(ypoint) - 2, max(ypoint) + 2)
    plt.xlim(-1, len(xpoint) + 2)

    plt.xlabel('采集频率(次)', fontsize=16)
    plt.ylabel("mems占用率%", fontsize=16)
    plt.title("APP进程MEM占用率%", fontsize=24)

    plt.legend()

    path = os.path.join(ScreenShot_DIR, 'mem_' + time_now)
    plt.savefig(path)
    
    
def mapping_batt():
    batts = excel.get_batts()
    batts = list(map(float, batts))
    times = [i for i in range(1, excel.times + 1)]
    # Draw graphics based on data
    plt.figure(figsize=(11, 7), dpi=600)
    # Generate grid
    plt.grid(axis="y")
    
    xpoint = np.array(times)
    ypoint = np.array(batts)
    
    plt.plot(xpoint, ypoint, color="red", linestyle="-", linewidth=1, label="手机电量")

    # Axis range
    plt.ylim(min(ypoint) - 2, max(ypoint) + 2)
    plt.xlim(-1, len(xpoint) + 2)

    plt.xlabel('采集频率(次)', fontsize=16)
    plt.ylabel("手机剩余batts%", fontsize=16)
    plt.title("手机剩余电量%", fontsize=24)

    plt.legend()

    path = os.path.join(ScreenShot_DIR, 'batt_' + time_now)
    plt.savefig(path)
    

def mapping_netflow():
    uploads, downloads, netflows = excel.get_netflows()
    uploads = list(map(float, uploads))
    downloads = list(map(float, downloads))
    netflows = list(map(float, netflows))
    # Draw graphics based on data
    plt.figure(figsize=(11, 7), dpi=600)
    # Generate grid
    plt.grid(axis="y")
    
    times = [i for i in range(1, excel.times + 1)]
    xpoint = np.array(times)
    
    ypoint1 = np.array(uploads)
    ypoint2 = np.array(downloads)
    ypoint3 = np.array(netflows)

    plt.plot(xpoint, ypoint1, color="blue", linestyle="-", marker='.',linewidth=1, label="上行流量")
    plt.plot(xpoint, ypoint2, color="green", linestyle="-",marker='.', linewidth=1, label="下行流量")
    plt.plot(xpoint, ypoint3, color="red", linestyle="-",marker='.', linewidth=1, label="总流量")
    
    # Axis range
    plt.ylim(min(ypoint3) - 1000, max(ypoint3) + 1000)
    plt.xlim(-1, len(xpoint) + 2)

    plt.xlabel('采集频率(次)', fontsize=16)
    plt.ylabel("消耗流量K", fontsize=16)
    plt.title("流量统计折线", fontsize=24)

    plt.legend()

    path = os.path.join(ScreenShot_DIR, 'netflow_' + time_now)
    plt.savefig(path)
    # plt,show()


def mapping_pic():
    # get datas
    cpus = excel.get_cpus()
    uploads, downloads, netflows = excel.get_netflows()
    batts = excel.get_batts()
    mems = excel.get_mems()
    # init datas
    cpus = list(map(float, cpus))
    mems = list(map(float, mems))
    batts = list(map(float, batts))
    uploads = list(map(float, uploads))
    downloads = list(map(float, downloads))
    netflows = list(map(float, netflows))
    # collection count
    times = [i for i in range(1, excel.times + 1)]
    
    # Draw graphics based on data
    plt.figure(figsize=(11, 7), dpi=600)
    # Generate grid
    plt.grid(axis="y")
    
    # xpoint of all
    xpoint = np.array(times)
    
    # cpu
    ypoint = np.array(cpus)
    plt.subplot(2, 2, 1)
    plt.plot(xpoint,ypoint)
    plt.title("cpu使用率%")
    
    # mem
    ypoint = np.array(mems)
    plt.subplot(2, 2, 2)
    plt.plot(xpoint,ypoint)
    plt.title("内存使用率%")
    
    # batt
    ypoint = np.array(batts)
    plt.subplot(2, 2, 4)
    plt.plot(xpoint,ypoint)
    plt.title("剩余电量%")
    
    # netflow   
    ypoint1 = np.array(uploads)
    ypoint2 = np.array(downloads)
    ypoint3 = np.array(netflows)
    plt.subplot(2, 2, 3)
    plt.plot(xpoint,ypoint1)
    plt.plot(xpoint,ypoint2)
    plt.plot(xpoint,ypoint3)
    plt.title("流量统计K")
    # main title
    plt.suptitle("Android应用性能指标数据")
    # display pic
    plt.legend() 
    path = os.path.join(ScreenShot_DIR, 'cpu_mem_netflow_' + time_now)
    # save pic
    plt.savefig(path)
    # plt.show()